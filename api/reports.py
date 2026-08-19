"""
Reports module — reworks Location/views.py's inventory_report/export_pdf/export_excel.

Fixes applied vs. the old code (data-breaking bugs, not stylistic changes — see
the migration review for detail):
  - "order" report: the old view had `return` INSIDE the row-building for loop,
    so it only ever returned a report with one row. Fixed to build all rows
    first, then return once.
  - "sales" report: filtered `status="completed"` (lowercase) against a model
    whose actual choices are capitalized ("Completed"), so it always matched
    zero orders. Fixed to compare against the real choice value.
  - "supplier" report: aggregated `Order.objects(supplier=s)`, but Order has no
    `supplier` field — replaced with the real relationship instead
    (Product.supplier), per your direction: supplier + their registered
    products/stock value.
  - "expiry" report: filtered on Product.expiry_date / Product.reorder_level,
    neither of which exist on Product — renamed to a Low-Stock report using
    the one part of the old logic that had real data behind it
    (quantity_available <= threshold), per your direction.
  - PDF/Excel export: the old export_pdf/export_excel always produced the same
    "top 5 products" summary regardless of which report was on screen. Now
    export renders the exact same headers/rows as the on-screen report, via
    one shared data builder both the JSON view and the export views call.
"""
import io
import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors as pdf_colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
from rest_framework.views import APIView
from rest_framework.response import Response
from .permissions import IsReportsStaff
from rest_framework import status as http_status
from django.http import HttpResponse

# Brand palette (matches react-web's index.css --color-primary tokens) so the
# exported documents read as the same product as the app, not a generic report.
BRAND_PRIMARY = pdf_colors.HexColor('#7986CB')
BRAND_DARK = pdf_colors.HexColor('#303F9F')
BRAND_TEXT = pdf_colors.HexColor('#0F172A')
BRAND_TEXT_SECONDARY = pdf_colors.HexColor('#64748B')
BRAND_BORDER = pdf_colors.HexColor('#E2E8F0')
BRAND_ZEBRA = pdf_colors.HexColor('#F3F4F8')

XL_PRIMARY = 'FF7986CB'
XL_DARK = 'FF303F9F'
XL_ZEBRA = 'FFF3F4F8'
XL_BORDER = 'FFE2E8F0'
XL_TEXT_INVERSE = 'FFFFFFFF'

_NUMERIC_HEADER_HINTS = ('value', 'amount', 'qty', 'total', 'paid', 'balance', 'owed', 'refunded', 'price', 'quantity', 'count')


def _is_numeric_column(header):
    h = header.lower()
    return any(hint in h for hint in _NUMERIC_HEADER_HINTS)

from product_Items.models import Product
from Orders.models import Order
from Location.models import Warehouse
from mysite.models import Supplier
from mysite.grn_models import GRN
from .po_models import PurchaseOrder, POPaymentRecord

LOW_STOCK_DEFAULT_THRESHOLD = 10


def _report_inventory(params):
    # One row per (product, warehouse) it actually has stock in — a product
    # received into two warehouses gets two rows, one per warehouse, not one
    # row with a single merged quantity. A product with no stock yet (still
    # just a catalog reference, no GRN confirmed) gets one row showing that.
    products = Product.objects.all()
    headers = ['Product', 'SKU', 'Category', 'Subcategory', 'UOM', 'Quantity', 'Warehouse', 'Created At']
    rows = []
    for p in products:
        base = [
            p.name, p.sku,
            p.category.name if p.category else '—',
            p.subcategory.name if p.subcategory else '—',
            p.uom.name if p.uom else '—',
        ]
        created = p.created_at.strftime('%d %b %Y, %I:%M %p') if p.created_at else '—'
        if not p.stock:
            rows.append(base + [0, '—', created])
        else:
            for line in p.stock:
                if not line.warehouse:
                    continue
                rows.append(base + [line.quantity_available, line.warehouse.warehouse_name, created])
    return {'title': 'Inventory Report', 'headers': headers, 'rows': rows}


def _report_orders(params):
    orders = Order.objects.all().order_by('-order_date')
    headers = ['Order ID', 'Customer', 'Order Date', 'Total Amount', 'Status', 'Payment Status']
    rows = [
        [
            str(o.id), o.customer_name,
            o.order_date.strftime('%d %b %Y') if o.order_date else '—',
            f'Rs. {o.total_amount:.2f}' if o.total_amount else 'Rs. 0.00',
            o.status, o.payment_status,
        ]
        for o in orders
    ]
    return {'title': 'Orders Report', 'headers': headers, 'rows': rows}


def _report_suppliers(params):
    suppliers = Supplier.objects.all()
    # "Products Supplied" lists the actual products (not just a count) so
    # this report answers "what did we buy from this supplier", tracked via
    # Product.supplier — kept in sync whenever a supplier GRN is confirmed
    # (see GRNConfirmView). That field only ever holds the *most recent*
    # supplier though, so it can't answer "how much have we purchased from
    # them in total" if a product was ever re-sourced from someone else —
    # "Total Qty Purchased" / "Total Purchase Value" are computed straight
    # from confirmed GRN line items instead, the actual receiving record,
    # so they stay accurate regardless of what Product.supplier says today.
    headers = [
        'Supplier', 'Company', 'Email', 'Phone', 'Products Supplied', 'Product Count',
        'Total Qty Purchased', 'Total Purchase Value', 'Current Stock Value',
    ]
    rows = []
    for s in suppliers:
        products = list(Product.objects(supplier=s))
        current_stock_value = sum(p.total_quantity * (p.price_per_unit or 0) for p in products)
        product_names = ', '.join(p.name for p in products) if products else '—'

        qty_purchased = 0
        value_purchased = 0.0
        for grn in GRN.objects(supplier=s, status='Confirmed'):
            for item in grn.items:
                qty_purchased += item.received_qty
                value_purchased += item.line_total

        rows.append([
            s.supplier_name, s.company_name, s.email, s.phone,
            product_names, len(products),
            qty_purchased, f'Rs. {value_purchased:.2f}', f'Rs. {current_stock_value:.2f}',
        ])
    return {'title': 'Supplier Report', 'headers': headers, 'rows': rows}


def _report_warehouses(params):
    warehouses = list(Warehouse.objects.all())
    headers = ['Warehouse', 'City', 'Total Products', 'Total Quantity', 'Low Stock Lines']
    rows = []
    row_meta = []
    for w in warehouses:
        # Only count each product's line *for this warehouse* — a product
        # also stocked elsewhere shouldn't inflate this warehouse's total.
        products = Product.objects(stock__warehouse=w)
        total_qty = 0
        product_count = 0
        low_count = 0
        for p in products:
            line = p.get_stock(w)
            if line:
                product_count += 1
                total_qty += line.quantity_available or 0
                if (line.quantity_available or 0) <= LOW_STOCK_DEFAULT_THRESHOLD:
                    low_count += 1
        rows.append([w.warehouse_name, w.city, product_count, total_qty, low_count])
        row_meta.append({'warehouse_id': str(w.id)})
    return {'title': 'Warehouse Report', 'headers': headers, 'rows': rows, 'row_meta': row_meta}


def _warehouse_stock_detail(warehouse_id):
    warehouse = Warehouse.objects(id=warehouse_id).first()
    if not warehouse:
        return None
    rows = []
    for p in Product.objects(stock__warehouse=warehouse):
        line = p.get_stock(warehouse)
        if not line:
            continue
        qty = line.quantity_available or 0
        rows.append({
            'product': p.name,
            'sku': p.sku,
            'quantity_available': qty,
            'price_per_unit': line.price_per_unit if line.price_per_unit is not None else p.price_per_unit,
            'is_low': qty <= LOW_STOCK_DEFAULT_THRESHOLD,
        })
    rows.sort(key=lambda r: r['quantity_available'])
    return {
        'warehouse': {'id': str(warehouse.id), 'name': warehouse.warehouse_name, 'city': warehouse.city},
        'threshold': LOW_STOCK_DEFAULT_THRESHOLD,
        'items': rows,
    }


def _report_sales(params):
    orders = Order.objects(status='Completed')
    headers = ['Order ID', 'Customer', 'Date', 'Total Amount', 'Payment Status']
    rows = [
        [
            str(o.id), o.customer_name,
            o.order_date.strftime('%d %b %Y') if o.order_date else '—',
            f'Rs. {o.total_amount:.2f}' if o.total_amount else 'Rs. 0.00',
            o.payment_status,
        ]
        for o in orders
    ]
    total_revenue = sum(o.total_amount for o in orders if o.total_amount)
    return {'title': 'Sales Report', 'headers': headers, 'rows': rows, 'summary': f'Total Revenue: Rs. {total_revenue:.2f}'}


def _report_low_stock(params):
    try:
        threshold = int(params.get('threshold', LOW_STOCK_DEFAULT_THRESHOLD))
    except (TypeError, ValueError):
        threshold = LOW_STOCK_DEFAULT_THRESHOLD
    # Low stock is a per-warehouse condition now — a product might be
    # plentiful in one warehouse and nearly out in another, so each
    # (product, warehouse) line is checked against the threshold on its own.
    headers = ['Product', 'SKU', 'UOM', 'Quantity Available', 'Warehouse']
    rows = []
    for p in Product.objects.all():
        for line in p.stock:
            if line.warehouse and (line.quantity_available or 0) <= threshold:
                rows.append([p.name, p.sku, p.uom.name if p.uom else '—', line.quantity_available, line.warehouse.warehouse_name])
    return {'title': f'Low-Stock Report (≤ {threshold})', 'headers': headers, 'rows': rows}


def _po_item_lines(po):
    """Per-PO line-item breakdown — shared by the export views' 'Details' tab/sheet."""
    headers = ['Product', 'Ordered Qty', 'Confirmed Price', 'Received Qty', 'Loss Qty', 'Line Total (Received)']
    rows = [
        [
            i.product.name if i.product else '—', i.ordered_qty,
            f'Rs. {(i.confirmed_price or i.proposed_price):.2f}', i.received_qty, i.loss_qty,
            f'Rs. {i.line_total_received:.2f}',
        ]
        for i in po.items
    ]
    return {'po_number': po.po_number, 'supplier': po.supplier.supplier_name if po.supplier else '—', 'headers': headers, 'rows': rows}


def _report_purchase_orders(params):
    pos = list(PurchaseOrder.objects.all())
    headers = ['PO Number', 'Supplier', 'Status', 'Ordered Value', 'Received Value', 'Loss Qty', 'Created']
    rows = []
    row_meta = []
    ordered_total = received_total = 0.0
    for po in pos:
        ordered_value = sum(i.ordered_qty * (i.confirmed_price or i.proposed_price) for i in po.items)
        received_value = po.received_total
        loss_qty = sum(i.loss_qty for i in po.items)
        ordered_total += ordered_value
        received_total += received_value
        rows.append([
            po.po_number, po.supplier.supplier_name, po.status,
            f'Rs. {ordered_value:.2f}', f'Rs. {received_value:.2f}', loss_qty,
            po.created_at.strftime('%d %b %Y') if po.created_at else '—',
        ])
        row_meta.append({'po_id': str(po.id)})
    loss_value = ordered_total - received_total
    return {
        'title': 'Purchase Orders Report', 'headers': headers, 'rows': rows, 'row_meta': row_meta,
        'summary': f'Ordered: Rs. {ordered_total:.2f}  |  Received: Rs. {received_total:.2f}  |  Loss: Rs. {loss_value:.2f}',
        'details': [_po_item_lines(po) for po in pos],
    }


def _report_po_payments(params):
    records = list(POPaymentRecord.objects.all())
    headers = ['PO Number', 'Supplier', 'Total Owed', 'Paid', 'Refunded', 'Balance', 'Status']
    rows = []
    row_meta = []
    total_paid = total_outstanding = total_refunded = 0.0
    for r in records:
        paid = sum(h.amount for h in r.history if h.kind == 'Payment')
        refunded = sum(h.amount for h in r.history if h.kind == 'Refund')
        total_paid += paid
        total_refunded += refunded
        total_outstanding += r.balance
        rows.append([
            r.po.po_number, r.po.supplier.supplier_name,
            f'Rs. {r.total_amount:.2f}', f'Rs. {paid:.2f}', f'Rs. {refunded:.2f}',
            f'Rs. {r.balance:.2f}', r.status,
        ])
        row_meta.append({'po_id': str(r.po.id)})
    return {
        'title': 'Supplier Payments Report', 'headers': headers, 'rows': rows, 'row_meta': row_meta,
        'summary': f'Paid: Rs. {total_paid:.2f}  |  Outstanding: Rs. {total_outstanding:.2f}  |  Refunded: Rs. {total_refunded:.2f}',
        'details': [_po_item_lines(r.po) for r in records],
    }


REPORT_BUILDERS = {
    'inventory': _report_inventory,
    'orders': _report_orders,
    'suppliers': _report_suppliers,
    'warehouses': _report_warehouses,
    'sales': _report_sales,
    'low-stock': _report_low_stock,
    'purchase-orders': _report_purchase_orders,
    'po-payments': _report_po_payments,
}


class ReportDataView(APIView):
    permission_classes = [IsReportsStaff]

    def get(self, request, report_type):
        builder = REPORT_BUILDERS.get(report_type)
        if not builder:
            return Response({'detail': f'Unknown report type "{report_type}".'}, status=http_status.HTTP_404_NOT_FOUND)
        return Response(builder(request.query_params))


class WarehouseStockDetailView(APIView):
    """Per-warehouse product/stock breakdown — the Warehouses report's row-click modal."""
    permission_classes = [IsReportsStaff]

    def get(self, request, warehouse_id):
        detail = _warehouse_stock_detail(warehouse_id)
        if not detail:
            return Response({'detail': 'Not found.'}, status=http_status.HTTP_404_NOT_FOUND)
        return Response(detail)


def _pdf_data_table(headers, rows, doc_width):
    """A zebra-striped, brand-colored table with numeric columns right-aligned."""
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('cell', parent=styles['Normal'], fontSize=8, textColor=BRAND_TEXT, leading=10)
    cell_style_r = ParagraphStyle('cell_r', parent=cell_style, alignment=TA_RIGHT)
    header_style = ParagraphStyle('header', parent=styles['Normal'], fontSize=8, textColor=pdf_colors.whitesmoke, fontName='Helvetica-Bold', leading=10)
    header_style_r = ParagraphStyle('header_r', parent=header_style, alignment=TA_RIGHT)

    numeric_cols = {i for i, h in enumerate(headers) if _is_numeric_column(h)}
    header_row = [Paragraph(h, header_style_r if i in numeric_cols else header_style) for i, h in enumerate(headers)]
    body = [
        [Paragraph(str(cell), cell_style_r if i in numeric_cols else cell_style) for i, cell in enumerate(row)]
        for row in rows
    ]
    table_data = [header_row] + body

    col_width = doc_width / max(len(headers), 1)
    table = Table(table_data, hAlign='LEFT', colWidths=[col_width] * len(headers), repeatRows=1)

    style = [
        ('BACKGROUND', (0, 0), (-1, 0), BRAND_DARK),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('LINEBELOW', (0, 0), (-1, 0), 1, BRAND_DARK),
        ('LINEBELOW', (0, -1), (-1, -1), 0.5, BRAND_BORDER),
        ('LINEBELOW', (0, 1), (-1, -2), 0.5, BRAND_BORDER),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]
    for i in range(1, len(table_data)):
        if i % 2 == 0:
            style.append(('BACKGROUND', (0, i), (-1, i), BRAND_ZEBRA))
    table.setStyle(TableStyle(style))
    return table


def _pdf_header(story, title, subtitle=None):
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle('brand', parent=styles['Normal'], fontSize=11, textColor=BRAND_DARK, fontName='Helvetica-Bold')
    title_style = ParagraphStyle('title', parent=styles['Title'], textColor=BRAND_TEXT, fontSize=20, spaceAfter=2)
    meta_style = ParagraphStyle('meta', parent=styles['Normal'], fontSize=9, textColor=BRAND_TEXT_SECONDARY)

    story.append(Paragraph('PRIMEFRESH SCM', brand_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(title, title_style))
    if subtitle:
        story.append(Paragraph(subtitle, meta_style))
    story.append(Paragraph(f"Generated {datetime.datetime.now().strftime('%d %B %Y, %I:%M %p')}", meta_style))
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width='100%', thickness=2, color=BRAND_PRIMARY, spaceAfter=14))


class ReportExportPdfView(APIView):
    permission_classes = [IsReportsStaff]

    def get(self, request, report_type):
        builder = REPORT_BUILDERS.get(report_type)
        if not builder:
            return Response({'detail': f'Unknown report type "{report_type}".'}, status=http_status.HTTP_404_NOT_FOUND)
        data = builder(request.query_params)
        preview = request.query_params.get('preview') == '1'

        response = HttpResponse(content_type='application/pdf')
        disposition = 'inline' if preview else 'attachment'
        response['Content-Disposition'] = f'{disposition}; filename="{report_type}_report.pdf"'

        doc = SimpleDocTemplate(response, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        story = []
        _pdf_header(story, data['title'], 'Summary')
        if data.get('summary'):
            summary_style = ParagraphStyle('summary', parent=styles['Normal'], fontSize=10, textColor=BRAND_DARK, fontName='Helvetica-Bold')
            story.append(Paragraph(data['summary'], summary_style))
            story.append(Spacer(1, 12))

        story.append(_pdf_data_table(data['headers'], data['rows'], doc.width))

        # Second tab (as a distinct section, since PDF has no literal tabs):
        # per-PO product line items, only present for PO-driven reports.
        details = data.get('details')
        if details:
            story.append(PageBreak())
            _pdf_header(story, data['title'], 'Details — products per Purchase Order')
            detail_heading = ParagraphStyle('detail_heading', parent=styles['Heading3'], textColor=BRAND_DARK, fontSize=12, spaceBefore=10, spaceAfter=6)
            for d in details:
                if not d['rows']:
                    continue
                story.append(Paragraph(f"{d['po_number']} &nbsp;-&nbsp; {d['supplier']}", detail_heading))
                story.append(_pdf_data_table(d['headers'], d['rows'], doc.width))
                story.append(Spacer(1, 10))

        doc.build(story)
        return response


def _xl_header_fill():
    return PatternFill(start_color=XL_DARK, end_color=XL_DARK, fill_type='solid')


def _xl_thin_border():
    side = Side(style='thin', color=XL_BORDER)
    return Border(left=side, right=side, top=side, bottom=side)


def _write_excel_table(ws, headers, rows, start_row=1):
    """Writes a styled table (header fill, zebra rows, borders, autofit, right-aligned numerics) starting at start_row. Returns the next free row."""
    border = _xl_thin_border()
    numeric_cols = {i for i, h in enumerate(headers) if _is_numeric_column(h)}

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.font = Font(bold=True, color=XL_TEXT_INVERSE, size=10)
        cell.fill = _xl_header_fill()
        cell.alignment = Alignment(horizontal='right' if (col - 1) in numeric_cols else 'left', vertical='center')
        cell.border = border
    ws.row_dimensions[start_row].height = 20

    for r, row in enumerate(rows):
        excel_row = start_row + 1 + r
        for col, cell_value in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=col, value=cell_value)
            cell.border = border
            cell.alignment = Alignment(horizontal='right' if (col - 1) in numeric_cols else 'left', vertical='center')
            cell.font = Font(size=10)
            if r % 2 == 1:
                cell.fill = PatternFill(start_color=XL_ZEBRA, end_color=XL_ZEBRA, fill_type='solid')

    # Autofit-ish: size each column to its longest cell (headers included), capped for readability.
    for col, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(row[col - 1])) for row in rows]) if rows else len(str(h))
        ws.column_dimensions[get_column_letter(col)].width = min(max(longest + 3, 12), 44)

    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1 + len(rows)


class ReportExportExcelView(APIView):
    permission_classes = [IsReportsStaff]

    def get(self, request, report_type):
        builder = REPORT_BUILDERS.get(report_type)
        if not builder:
            return Response({'detail': f'Unknown report type "{report_type}".'}, status=http_status.HTTP_404_NOT_FOUND)
        data = builder(request.query_params)

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{report_type}_report.xlsx"'

        wb = openpyxl.Workbook()
        summary_ws = wb.active
        summary_ws.title = 'Summary'

        title_cell = summary_ws.cell(row=1, column=1, value=data['title'])
        title_cell.font = Font(bold=True, size=14, color=XL_DARK)
        generated_cell = summary_ws.cell(row=2, column=1, value=f"Generated {datetime.datetime.now().strftime('%d %B %Y, %I:%M %p')}")
        generated_cell.font = Font(size=9, italic=True, color='FF64748B')
        next_row = 3
        if data.get('summary'):
            summary_cell = summary_ws.cell(row=3, column=1, value=data['summary'])
            summary_cell.font = Font(bold=True, size=10, color=XL_DARK)
            next_row = 4
        next_row += 1

        _write_excel_table(summary_ws, data['headers'], data['rows'], start_row=next_row)

        # Second tab: per-PO product line items, only for PO-driven reports.
        details = data.get('details')
        if details:
            detail_ws = wb.create_sheet('Details')
            row = 1
            any_written = False
            for d in details:
                if not d['rows']:
                    continue
                any_written = True
                po_cell = detail_ws.cell(row=row, column=1, value=f"{d['po_number']}  -  {d['supplier']}")
                po_cell.font = Font(bold=True, size=11, color=XL_DARK)
                row += 1
                row = _write_excel_table(detail_ws, d['headers'], d['rows'], start_row=row)
                row += 2
            if not any_written:
                detail_ws.cell(row=1, column=1, value='No purchase order line items yet.')

        wb.save(response)
        return response
